"""Generate our reconstruction plan and apply consistent pagination to drafts."""
from pathlib import Path
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def main(root):
    import importlib.metadata
    print('Document Python:', sys.version, 'python-docx:', importlib.metadata.version('python-docx'))
    paper = root/'results/paper'
    d = Document()
    d.add_paragraph('Paper reorganisation and reproduction plan', 'Title')
    d.add_paragraph('This working plan maps the advisor review to independently computed project results. It is a reconstruction plan, not evidence that all inherited scientific interpretations are established.')
    sections = [
        ('Data and provenance', 'Use review_package/data/combined_E0_final.csv and combined_R0c_final.csv. Retain exposure records, including zero affected customers. For final recovery regressions exclude zero-customer records and report this change explicitly. Keep incident-conditional analyses separate from district-day occurrence probabilities.'),
        ('Extended paper', 'Organise the extended paper as Introduction, Data, Methods, Results, Discussion and Supplementary Information. Methods cover outcome construction, the 12-specification ladder, profile knot search, district bootstrap, spatial and temporal validation, two-way clustered inference, and district-day fragility. Results present pooled response shapes, weather-attributed comparisons, variance decomposition, district-day occurrence and temporal checks.'),
        ('Short paper', 'Focus the Nature Communications extract on the operational district-day fragility estimate, supported by the estimated exposure response and recovery contrast. Move model-search details, ordinal and negative-binomial experiments, calibration panels and full coefficient tables to supplementary material.'),
        ('Figures and tables', 'Reuse the study map from our own docs/Draft.docx. Generate fig1 through fig14 and figP1/figP2 from this run. Populate final-model, weather-only, predicted-versus-observed and fragility tables from computed CSV/JSON files. The numbered series intentionally has no fig7, as in the supplied analysis code.'),
        ('Interpretation checks', 'Describe the low-wind slope using its estimated sign and uncertainty, rather than assuming it is flat. The E0 plateau is a constraint on the main gust effect at mean pressure; the pressure interaction remains. A non-rejected restriction is not proof of a physical law. The fitted background parameter is not an observed technical-fault rate. Interpolated weather on days without events remains a proxy. Avoid calling the fixed-specification temporal split independent confirmation of knot selection.'),
        ('Outstanding manuscript work', 'The inherited long and short drafts retain substantial hand-entered narrative. Before submission, update all narrative numbers from COMPUTED_RESULTS.md and the source CSV/JSON tables; review causal explanations, novelty claims and references. Exact sentence reproduction is distinct from scientific verification.'),
        ('Code and output scope', 'main_new.py controls analysis_new/. Each run stores its purpose, settings, status, source hashes and input hashes in results/new/<timestamp>/run.json. Results are in that run’s results/model_selection, final_models, weather_only, figures and paper directories. See docs/NEW_ANALYSIS_GUIDE.md and LOG.md for the dated implementation history.')]
    for title, body in sections:
        d.add_heading(title, level=1); d.add_paragraph(body)
    d.save(paper/'Paper_reorganisation_plan.docx')
    # The source templates have wide tables. Fit them to the text width and
    # prevent split rows; keep figures with their captions and titles black.
    for path in [root/'results/Model_exploration_predicted_vs_observed.docx', *paper.glob('*.docx')]:
        doc = Document(path)
        for sec in doc.sections:
            sec.page_width = Inches(8.5); sec.page_height = Inches(11)
            sec.left_margin = Inches(.65); sec.right_margin = Inches(.65)
        title = doc.styles['Title']; title.font.color.rgb = RGBColor(0, 0, 0)
        # Some Word templates give Title a theme-coloured paragraph border.
        for element in [doc._element, doc.styles.element]:
            for border in element.xpath('.//w:pBdr'):
                border.getparent().remove(border)
        for p in doc.paragraphs:
            if p.style is not None and (p.style.name == 'Title' or p.style.name.lower().startswith('heading')):
                p.paragraph_format.keep_with_next = True
            if p.text.startswith('Table ') and len(p.text) < 450:
                previous = p._p.getprevious()
                below_table = previous is not None and previous.tag == qn('w:tbl')
                p.paragraph_format.keep_with_next = not below_table
                if below_table:
                    # A caption below a table belongs to its final row, not to
                    # the next large image (which would move it to a new page).
                    last_row = previous.findall(qn('w:tr'))[-1]
                    for cell in last_row.findall(qn('w:tc')):
                        last_p = cell.findall(qn('w:p'))[-1]
                        props = last_p.find(qn('w:pPr'))
                        if props is None:
                            props = OxmlElement('w:pPr'); last_p.insert(0, props)
                        if props.find(qn('w:keepNext')) is None:
                            props.append(OxmlElement('w:keepNext'))
            if p._p.xpath('.//w:drawing'):
                p.paragraph_format.keep_with_next = True
        for table in doc.tables:
            table.autofit = False
            widths = [int(c.get(qn('w:w'))) for c in table._tbl.tblGrid.gridCol_lst]
            target = int(7.2 * 1440)
            scaled = [round(w * target / sum(widths)) for w in widths]
            for c, width in zip(table._tbl.tblGrid.gridCol_lst, scaled):
                c.set(qn('w:w'), str(width))
            table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:w'), str(target))
            for row in table.rows:
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for cell, width in zip(row.cells, scaled):
                    cell.width = Inches(width/1440)
        doc.save(path)


if __name__ == '__main__':
    main(Path(sys.argv[1]))
